from __future__ import annotations

import json
import os
import shlex
import socket
import subprocess
import time
import urllib.error
import urllib.request
import zipfile
from abc import ABC, abstractmethod
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Any

from backend.config import AppConfig, get_settings
from backend.services.command_runner import resolve_command
from backend.services.security import mask_sensitive


class AgentAdapter(ABC):
    @abstractmethod
    def health_check(self) -> dict[str, Any]:
        pass

    @abstractmethod
    def run_task(
        self,
        task_name: str,
        prompt: str,
        workspace: Path,
        extra_context: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        pass


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _timestamp() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _headers(settings: AppConfig) -> dict[str, str]:
    headers = {"Content-Type": "application/json"}
    api_key = os.getenv(settings.hermes_api_key_env, "")
    if api_key:
        headers["Authorization"] = f"Bearer {api_key}"
    return headers


def _join_url(endpoint: str, path: str) -> str:
    return endpoint.rstrip("/") + "/" + path.lstrip("/")


class HermesAPIAdapter(AgentAdapter):
    def __init__(self, settings: AppConfig | None = None):
        self.settings = settings or get_settings()

    def health_check(self) -> dict[str, Any]:
        if not self.settings.hermes_endpoint:
            return {
                "name": "hermes",
                "ok": False,
                "status": "missing_endpoint",
                "detail": "HERMES_ENDPOINT is not configured.",
            }

        started = time.perf_counter()
        request = urllib.request.Request(
            _join_url(self.settings.hermes_endpoint, self.settings.hermes_health_path),
            headers=_headers(self.settings),
            method="GET",
        )
        try:
            with urllib.request.urlopen(request, timeout=10) as response:
                raw = response.read().decode("utf-8", errors="replace")
            return {
                "name": "hermes",
                "ok": True,
                "status": "success",
                "detail": "Hermes API health check passed.",
                "raw_output": mask_sensitive(raw),
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
            }
        except urllib.error.HTTPError as exc:
            error_code = "hermes_auth_failed" if exc.code in {401, 403} else "hermes_failed"
            return {
                "name": "hermes",
                "ok": False,
                "status": "failed",
                "error_code": error_code,
                "detail": "Hermes API health check failed.",
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
            }
        except (TimeoutError, socket.timeout):
            return {
                "name": "hermes",
                "ok": False,
                "status": "failed",
                "error_code": "hermes_timeout",
                "detail": "Hermes API health check timed out.",
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
            }
        except (urllib.error.URLError, OSError):
            return {
                "name": "hermes",
                "ok": False,
                "status": "failed",
                "error_code": "hermes_bridge_unavailable",
                "detail": "Hermes API bridge is unavailable.",
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
            }

    def run_task(
        self,
        task_name: str,
        prompt: str,
        workspace: Path,
        extra_context: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        if not self.settings.hermes_endpoint:
            return {
                "status": "failed",
                "task_name": task_name,
                "result": None,
                "error_code": "hermes_bridge_unavailable",
                "error": "Hermes endpoint is not configured.",
                "elapsed_ms": 0,
                "log_file": None,
            }

        payload = {
            "task_name": task_name,
            "prompt": prompt,
            "workspace": str(workspace),
            "extra_context": extra_context or {},
        }
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        request = urllib.request.Request(
            _join_url(self.settings.hermes_endpoint, self.settings.hermes_run_path),
            data=body,
            headers=_headers(self.settings),
            method="POST",
        )
        started = time.perf_counter()
        try:
            with urllib.request.urlopen(request, timeout=300) as response:
                raw = response.read().decode("utf-8", errors="replace")
            parsed = json.loads(raw) if raw else {}
            if not isinstance(parsed, dict):
                raise ValueError("Hermes bridge response must be a JSON object.")
            log_path = workspace / "logs" / f"hermes_{task_name}.log"
            log_path.parent.mkdir(parents=True, exist_ok=True)
            log_path.write_text(mask_sensitive(raw), encoding="utf-8")
            return {
                "status": "success",
                "task_name": task_name,
                "result": parsed,
                "error_code": None,
                "error": None,
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
                "log_file": log_path.relative_to(workspace).as_posix(),
            }
        except urllib.error.HTTPError as exc:
            error_code = "hermes_auth_failed" if exc.code in {401, 403} else "hermes_failed"
            return {
                "status": "failed",
                "task_name": task_name,
                "result": None,
                "error_code": error_code,
                "error": "Hermes bridge returned an HTTP error.",
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
                "log_file": None,
            }
        except (TimeoutError, socket.timeout):
            return {
                "status": "failed",
                "task_name": task_name,
                "result": None,
                "error_code": "hermes_timeout",
                "error": "Hermes bridge request timed out.",
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
                "log_file": None,
            }
        except (json.JSONDecodeError, ValueError):
            return {
                "status": "failed",
                "task_name": task_name,
                "result": None,
                "error_code": "hermes_invalid_result",
                "error": "Hermes bridge returned an invalid result.",
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
                "log_file": None,
            }
        except (urllib.error.URLError, OSError):
            return {
                "status": "failed",
                "task_name": task_name,
                "result": None,
                "error_code": "hermes_bridge_unavailable",
                "error": "Hermes bridge is unavailable.",
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
                "log_file": None,
            }


class MockHermesAdapter(AgentAdapter):
    def health_check(self) -> dict[str, Any]:
        return {
            "name": "hermes",
            "ok": True,
            "status": "mock",
            "detail": "Hermes API is mocked. Configure HERMES_MODE=api and HERMES_ENDPOINT for live calls.",
        }

    def run_task(
        self,
        task_name: str,
        prompt: str,
        workspace: Path,
        extra_context: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        context = extra_context or {}
        request = context.get("request", {})
        title = request.get("title") or "未命名报告"
        domain = request.get("domain") or "通用"
        audience = request.get("audience") or "专业读者"
        language = request.get("language") or "中文"
        result: dict[str, Any]

        if task_name == "planner":
            result = {
                "task_understanding": f"围绕《{title}》生成面向{audience}的{domain}报告。",
                "outline": [
                    {"section": "背景与问题定义", "goal": "界定主题、应用场景和报告边界"},
                    {"section": "关键事实与证据", "goal": "整理公开资料、专利、论文和产业信息"},
                    {"section": "竞争格局与风险", "goal": "归纳技术路线、主要参与方和不确定性"},
                    {"section": "结论与建议", "goal": "形成可汇报的判断和行动建议"},
                ],
                "search_questions": [
                    f"{title} 最新产业进展",
                    f"{title} 专利 布局 竞争",
                    f"{title} 技术路线 风险",
                ],
                "figures_needed": ["技术路线对比表", "主要机构/企业矩阵", "风险与建议优先级图"],
                "risks": ["公开资料时间滞后", "专利族口径不一致", "行业宣传材料可能带有偏向"],
                "success_criteria": ["结构完整", "事实有来源", "PPT可汇报", "Word可追溯"],
                "language": language,
            }
        elif task_name == "writer":
            sources = read_json(workspace / "research" / "sources.json", [])
            result = {
                "content_markdown": (
                    f"# {title}\n\n"
                    f"## 背景与问题定义\n本报告面向{audience}，聚焦{domain}场景下的核心问题、公开证据和落地建议。\n\n"
                    "## 关键事实与证据\n当前版本使用结构化工作流沉淀来源，后续可替换为 Hermes API 输出的正式内容。\n\n"
                    "## 竞争格局与风险\n需要区分已验证事实、推断判断和待人工确认的信息。\n\n"
                    "## 结论与建议\n建议优先补齐高风险事实来源，再进入模板化排版与终审。"
                ),
                "slides_storyboard": [
                    {
                        "slide_no": 1,
                        "slide_type": "cover",
                        "title": title,
                        "bullets": [f"领域：{domain}", f"受众：{audience}"],
                        "citations": [],
                    },
                    {
                        "slide_no": 2,
                        "slide_type": "content",
                        "title": "研究框架",
                        "bullets": ["背景与问题定义", "关键事实与证据", "竞争格局与风险", "结论与建议"],
                        "citations": [],
                    },
                    {
                        "slide_no": 3,
                        "slide_type": "content",
                        "title": "证据与风险",
                        "bullets": [
                            f"已记录来源数量：{len(sources)}",
                            "事实、推断和待确认信息分层展示",
                            "审稿意见以 review.json 沉淀",
                        ],
                        "citations": [],
                    },
                ],
            }
        elif task_name in {"reviewer", "final_reviewer"}:
            round_id = int(context.get("round_id", 1))
            passed = task_name == "final_reviewer" or round_id > 1
            score = 90 if passed else 78
            result = {
                "score": score,
                "pass": passed,
                "blocking_issues": []
                if passed
                else [
                    {
                        "type": "traceability",
                        "severity": "high",
                        "location": "slide_3",
                        "comment": "需要明确标记哪些结论来自已保存 sources.json，哪些仍是待验证判断。",
                    }
                ],
                "minor_issues": [
                    {
                        "type": "layout",
                        "location": "output",
                        "comment": "正式接入模板后需要检查标题长度和图表占位符。",
                    }
                ],
                "revision_instruction": "优先补充事实来源标记，再进行版式修复。" if not passed else "通过终审。",
            }
        else:
            result = {"status": "success", "message": f"Mock Hermes handled {task_name}."}

        log_path = workspace / "logs" / f"hermes_{task_name}.log"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        log_path.write_text(json.dumps({"prompt": prompt, "result": result}, ensure_ascii=False, indent=2), encoding="utf-8")
        return {"status": "success", "task_name": task_name, "result": result, "log_file": str(log_path)}


class CodexCLIAdapter(AgentAdapter):
    def __init__(self, settings: AppConfig | None = None):
        self.settings = settings or get_settings()

    def health_check(self) -> dict[str, Any]:
        resolved = resolve_command(self.settings.codex_command)
        if not resolved:
            return {
                "name": "codex",
                "ok": False,
                "status": "missing",
                "command": self.settings.codex_command,
                "detail": "Codex command was not found on PATH.",
            }
        try:
            completed = subprocess.run(
                [resolved, "--version"],
                capture_output=True,
                text=True,
                timeout=10,
                shell=False,
                encoding="utf-8",
                errors="replace",
            )
            raw = mask_sensitive((completed.stdout or "") + (completed.stderr or "")).strip()
            return {
                "name": "codex",
                "ok": completed.returncode == 0,
                "status": "success" if completed.returncode == 0 else "failed",
                "command": resolved,
                "version": raw.splitlines()[0] if raw else None,
                "raw_output": raw,
                "detail": f"exit_code={completed.returncode}",
            }
        except (subprocess.TimeoutExpired, OSError) as exc:
            return {
                "name": "codex",
                "ok": False,
                "status": "failed",
                "command": resolved,
                "detail": str(exc),
            }

    def run_task(
        self,
        task_name: str,
        prompt: str,
        workspace: Path,
        extra_context: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        resolved = resolve_command(self.settings.codex_command)
        if not resolved:
            raise RuntimeError("Codex command was not found on PATH.")

        prompt_file = workspace / "prompts" / f"codex_{task_name}.md"
        prompt_file.parent.mkdir(parents=True, exist_ok=True)
        prompt_file.write_text(prompt, encoding="utf-8")
        args = shlex.split(self.settings.codex_exec_args)
        command = [resolved, *args]
        started = time.perf_counter()
        try:
            completed = subprocess.run(
                command,
                input=prompt,
                cwd=str(workspace),
                capture_output=True,
                text=True,
                timeout=1800,
                shell=False,
                encoding="utf-8",
                errors="replace",
            )
            raw = mask_sensitive((completed.stdout or "") + "\n\nSTDERR:\n" + (completed.stderr or ""))
            log_path = workspace / "logs" / f"codex_{task_name}.log"
            log_path.parent.mkdir(parents=True, exist_ok=True)
            log_path.write_text(raw, encoding="utf-8")
            return {
                "status": "success" if completed.returncode == 0 else "failed",
                "task_name": task_name,
                "returncode": completed.returncode,
                "stdout": mask_sensitive(completed.stdout or ""),
                "stderr": mask_sensitive(completed.stderr or ""),
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
                "log_file": str(log_path),
            }
        except (subprocess.TimeoutExpired, OSError) as exc:
            return {
                "status": "failed",
                "task_name": task_name,
                "error": str(exc),
                "elapsed_ms": int((time.perf_counter() - started) * 1000),
            }


class MockCodexAdapter(AgentAdapter):
    def health_check(self) -> dict[str, Any]:
        return {
            "name": "codex",
            "ok": True,
            "status": "mock",
            "detail": "Codex CLI is mocked. Configure CODEX_MODE=cli to call the local Codex login.",
        }

    def run_task(
        self,
        task_name: str,
        prompt: str,
        workspace: Path,
        extra_context: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        context = extra_context or {}
        request = context.get("request", {})
        files_changed: list[str] = []

        prompt_path = workspace / "prompts" / f"codex_{task_name}.md"
        prompt_path.parent.mkdir(parents=True, exist_ok=True)
        prompt_path.write_text(prompt, encoding="utf-8")

        if task_name == "builder":
            self._write_build_script(workspace)
            files_changed.extend(self._write_outputs(workspace, request))
            summary = "Mock Codex generated output files and build metadata."
        elif task_name == "fixer":
            patch_note = {
                "fixed_at": _timestamp(),
                "source_review": context.get("review_file"),
                "changes": ["标记事实来源状态", "补充模板接入前的版式检查提示"],
            }
            write_json(workspace / "build" / "revision_patch.json", patch_note)
            files_changed.append("build/revision_patch.json")
            summary = "Mock Codex applied review instructions."
        else:
            summary = f"Mock Codex handled {task_name}."

        result = {
            "status": "success",
            "files_changed": files_changed,
            "summary": summary,
            "tests": {
                "workspace_write_check": "passed",
                "artifact_manifest_check": "passed",
            },
        }
        write_json(workspace / "build" / f"codex_{task_name}_result.json", result)
        log_path = workspace / "logs" / f"codex_{task_name}.log"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        log_path.write_text(json.dumps({"prompt": prompt, "result": result}, ensure_ascii=False, indent=2), encoding="utf-8")
        return {"status": "success", "task_name": task_name, "result": result, "log_file": str(log_path)}

    def _write_build_script(self, workspace: Path) -> None:
        script = (
            "from pathlib import Path\n\n"
            "workspace = Path(__file__).resolve().parents[1]\n"
            "print(f'Build workspace: {workspace}')\n"
        )
        path = workspace / "build" / "generate_ppt.py"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(script, encoding="utf-8")

    def _write_outputs(self, workspace: Path, request: dict[str, Any]) -> list[str]:
        changed: list[str] = []
        task_type = request.get("task_type", "ppt")
        title = request.get("title") or "Untitled Report"
        storyboard = read_json(workspace / "draft" / "slides_storyboard.json", [])
        content = (workspace / "draft" / "content.md").read_text(encoding="utf-8") if (workspace / "draft" / "content.md").exists() else title

        if task_type in {"ppt", "ppt_word"}:
            ppt_path = workspace / "output" / "final_presentation.pptx"
            ppt_path.parent.mkdir(parents=True, exist_ok=True)
            try:
                from pptx import Presentation

                presentation = Presentation()
                for slide in storyboard or [{"title": title, "bullets": []}]:
                    layout = presentation.slide_layouts[0] if slide.get("slide_type") == "cover" else presentation.slide_layouts[1]
                    ppt_slide = presentation.slides.add_slide(layout)
                    ppt_slide.shapes.title.text = str(slide.get("title") or title)
                    if len(ppt_slide.placeholders) > 1:
                        body = ppt_slide.placeholders[1].text_frame
                        body.clear()
                        for bullet in slide.get("bullets", []):
                            paragraph = body.add_paragraph()
                            paragraph.text = str(bullet)
                            paragraph.level = 0
                presentation.save(ppt_path)
            except Exception:
                ppt_path.write_text("PPT placeholder generated because python-pptx is unavailable.", encoding="utf-8")
            changed.append("output/final_presentation.pptx")

        if task_type in {"word", "ppt_word"}:
            docx_path = workspace / "output" / "final_report.docx"
            docx_path.parent.mkdir(parents=True, exist_ok=True)
            try:
                from docx import Document

                document = Document()
                document.add_heading(title, level=1)
                for block in content.split("\n\n"):
                    text = block.strip()
                    if text.startswith("# "):
                        document.add_heading(text[2:], level=1)
                    elif text.startswith("## "):
                        document.add_heading(text[3:], level=2)
                    elif text:
                        document.add_paragraph(text)
                document.save(docx_path)
            except Exception:
                self._write_minimal_docx(docx_path, title, content)
            changed.append("output/final_report.docx")

        manifest = {
            "generated_at": _timestamp(),
            "files": changed,
            "mode": "mock_codex",
        }
        write_json(workspace / "output" / "manifest.json", manifest)
        changed.append("output/manifest.json")
        return changed

    def _write_minimal_docx(self, path: Path, title: str, content: str) -> None:
        paragraphs = [line.strip() for line in content.replace("\r\n", "\n").split("\n") if line.strip()]
        if not paragraphs:
            paragraphs = [title]
        document_body = "\n".join(
            f"<w:p><w:r><w:t>{escape(paragraph)}</w:t></w:r></w:p>"
            for paragraph in paragraphs
        )
        document_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body>{document_body}<w:sectPr /></w:body>"
            "</w:document>"
        )
        content_types = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>"
        )
        rels = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/>'
            "</Relationships>"
        )
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as docx:
            docx.writestr("[Content_Types].xml", content_types)
            docx.writestr("_rels/.rels", rels)
            docx.writestr("word/document.xml", document_xml)


class MockChromeAdapter(AgentAdapter):
    def health_check(self) -> dict[str, Any]:
        return {
            "name": "chrome",
            "ok": True,
            "status": "mock",
            "detail": "Browser search is mocked; sources.json will record planned queries.",
        }

    def run_task(
        self,
        task_name: str,
        prompt: str,
        workspace: Path,
        extra_context: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        queries = read_json(workspace / "research" / "search_queries.json", [])
        sources = [
            {
                "title": f"待检索来源：{query}",
                "url": "",
                "date": "",
                "claim_supported": "需要接入 Chrome/browser 工具后验证。",
                "used_in_section": "关键事实与证据",
                "status": "pending_verification",
            }
            for query in queries
        ]
        write_json(workspace / "research" / "sources.json", sources)
        log_path = workspace / "logs" / "browser.log"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        log_path.write_text(json.dumps({"queries": queries, "sources": sources}, ensure_ascii=False, indent=2), encoding="utf-8")
        return {
            "status": "success",
            "task_name": task_name,
            "result": {"sources": sources},
            "log_file": str(log_path),
        }


def make_hermes_adapter(settings: AppConfig | None = None) -> AgentAdapter:
    config = settings or get_settings()
    if config.hermes_mode == "api":
        return HermesAPIAdapter(config)
    return MockHermesAdapter()


def make_codex_adapter(settings: AppConfig | None = None) -> AgentAdapter:
    config = settings or get_settings()
    if config.codex_mode == "cli":
        return CodexCLIAdapter(config)
    return MockCodexAdapter()


def make_chatgpt_adapter(settings: AppConfig | None = None) -> AgentAdapter:
    from backend.services.chatgpt_adapter import ChatGPTAdapter, MockChatGPTAdapter

    config = settings or get_settings()
    if config.chatgpt_mode == "cdp":
        return ChatGPTAdapter(config)
    return MockChatGPTAdapter()
